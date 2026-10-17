"""Importar produtos por planilha — mounted under /api/products/import.

Fluxo: o ADMIN baixa o modelo (.xlsx), preenche (ou usa a planilha que já tem),
envia com dry_run=true para ver a prévia e depois envia de novo com
dry_run=false para gravar. O SKU é a chave: SKU novo cria o produto, SKU que
já existe atualiza só as colunas preenchidas. Linhas com erro são ignoradas.
"""

import csv
import io
import re
import unicodedata
from datetime import datetime, timezone
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pymongo import UpdateOne
from pydantic import BaseModel

from lib.auth import Principal, require, scoped
from lib.db import db
from lib.plans import limits_for_org
from models.inventory import Product, StockMovement
from models.org import PLAN_LABELS

router = APIRouter(prefix="/products/import", tags=["products"])

MAX_ROWS = 50000
MAX_BYTES = 25 * 1024 * 1024
BATCH = 1000
PREVIEW_ROWS = 200
DEFAULT_CATEGORY = "Geral"

# cabeçalho normalizado (sem acento, minúsculo, só letras/números) -> campo
HEADER_ALIASES: dict[str, str] = {}
for field, names in {
    "name": ["nome", "produto", "descricao", "nomedoproduto", "descricaodoproduto", "item"],
    "sku": ["sku", "codigo", "cod", "codigointerno", "referencia", "ref", "codigodoproduto", "codproduto"],
    "barcode": ["codigodebarras", "codbarras", "barras", "ean", "gtin", "ean13", "codigobarras"],
    "category": ["categoria", "grupo", "departamento", "linha", "secao"],
    "cost_price": ["custo", "precodecusto", "precocusto", "valorcusto", "valordecusto"],
    "sale_price": ["preco", "precodevenda", "precovenda", "valorvenda", "valordevenda", "venda", "valor"],
    "current_stock": ["estoque", "estoqueatual", "quantidade", "qtd", "qtde", "saldo", "saldoatual"],
    "min_stock": ["estoqueminimo", "minimo", "estmin", "estoquemin", "qtdminima"],
}.items():
    for alias in names:
        HEADER_ALIASES[alias] = field

TEMPLATE_HEADERS = [
    "Nome", "SKU", "Código de barras", "Categoria",
    "Preço de custo", "Preço de venda", "Estoque", "Estoque mínimo",
]


class ImportRow(BaseModel):
    row: int
    sku: str = ""
    name: str = ""
    action: str  # CRIAR | ATUALIZAR | SEM_MUDANCA | ERRO
    message: str = ""


class ImportResult(BaseModel):
    dry_run: bool
    total_rows: int
    created: int
    updated: int
    unchanged: int
    errors: int
    rows: List[ImportRow]
    blocked: str = ""  # motivo que impede a importação inteira (ex.: limite do plano)
    message: str


# ------------------------------------------------------------------ parsing
def _norm(text: object) -> str:
    raw = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", raw.lower())


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_number(value: object) -> Optional[float]:
    """Aceita 12,50 · 1.234,56 · 1,234.56 · R$ 10 · 12.5 — None se vazio."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if not text or text in {"-", ",", "."}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", text):
        text = text.replace(".", "")  # 1.500 = mil e quinhentos
    try:
        return float(text)
    except ValueError:
        raise ValueError(f"valor inválido: {value}") from None


def _read_rows(filename: str, content: bytes) -> list[list[object]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or content[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="Não consegui abrir a planilha. Salve como .xlsx ou .csv.")
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Formato .xls antigo: abra no Excel e salve como .xlsx.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(status_code=400, detail="Não consegui ler o arquivo CSV.")
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") >= first.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _map_headers(header: list[object]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, title in enumerate(header):
        field = HEADER_ALIASES.get(_norm(title))
        if field and field not in mapping:
            mapping[field] = idx
    return mapping


# ------------------------------------------------------------------ routes
@router.get("/template.xlsx")
async def download_template(principal: Principal = Depends(require("products:write"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
    ws.append(["Pastilha de freio dianteira", "PF-1020", "7891234567895", "Freios", 45.9, 89.9, 12, 3])
    ws.append(["Filtro de óleo", "FO-330", "", "Filtros", "12,50", "24,90", 30, 5])
    for col, width in zip("ABCDEFGH", (36, 14, 18, 16, 15, 15, 10, 15)):
        ws.column_dimensions[col].width = width
    help_ws = wb.create_sheet("Como preencher")
    for line in (
        ["Como preencher a planilha"],
        [""],
        ["Obrigatórios: Nome e SKU (o código interno do produto). Produto novo também precisa do Preço de venda."],
        ["SKU já cadastrado: o produto é ATUALIZADO só nas colunas preenchidas."],
        ["Estoque: em produto existente, a quantidade vira o novo estoque (fica registrado como ajuste)."],
        ["Preços podem ter vírgula (12,50) ou ponto (12.50). Categoria vazia vira 'Geral'."],
        ["Apague as 2 linhas de exemplo antes de importar. Até 50.000 linhas por arquivo."],
    ):
        help_ws.append(line)
    help_ws.column_dimensions["A"].width = 110
    help_ws["A1"].font = Font(bold=True, size=13)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo-produtos.xlsx"'},
    )


@router.post("", response_model=ImportResult)
async def import_products(
    file: UploadFile = File(...),
    dry_run: bool = Query(default=True),
    principal: Principal = Depends(require("products:write")),
):
    content = await file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Arquivo grande demais (máximo 25 MB).")
    if not content:
        raise HTTPException(status_code=400, detail="Arquivo vazio.")

    raw = _read_rows(file.filename or "", content)
    # primeira linha com algum cabeçalho reconhecido (pula títulos acima da tabela)
    header_idx, mapping = -1, {}
    for idx, candidate in enumerate(raw[:10]):
        mapping = _map_headers(candidate)
        if "name" in mapping and "sku" in mapping:
            header_idx = idx
            break
    if header_idx < 0:
        raise HTTPException(
            status_code=400,
            detail="Não achei as colunas Nome e SKU. Use o modelo de planilha ou renomeie os cabeçalhos.",
        )
    body = [
        (header_idx + 2 + offset, r)
        for offset, r in enumerate(raw[header_idx + 1 :])
        if any(_cell(v) for v in r)
    ]
    if len(body) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"Máximo de {MAX_ROWS} linhas por arquivo.")

    existing = {
        d["sku"]: d
        for d in await db.products.find(scoped(principal)).to_list(None)
    }
    barcode_owner = {d.get("barcode"): d["sku"] for d in existing.values() if d.get("barcode")}

    def get(r: list[object], field: str) -> object:
        idx = mapping.get(field)
        return r[idx] if idx is not None and idx < len(r) else None

    rows: list[ImportRow] = []
    plans: list[tuple[str, dict, Optional[dict]]] = []  # (action, data, existing_doc)
    seen_skus: set[str] = set()
    for line, r in body:
        sku = _cell(get(r, "sku")).upper()[:40]
        name = _cell(get(r, "name"))[:120]
        try:
            if not sku:
                raise ValueError("SKU vazio")
            if sku in seen_skus:
                raise ValueError("SKU repetido na planilha")
            seen_skus.add(sku)
            data: dict = {}
            if name:
                data["name"] = name
            category = _cell(get(r, "category"))[:60]
            if category:
                data["category"] = category
            barcode = re.sub(r"\s+", "", _cell(get(r, "barcode")))[:40]
            if barcode:
                owner = barcode_owner.get(barcode)
                if owner and owner != sku:
                    raise ValueError(f"código de barras já usado no SKU {owner}")
                data["barcode"] = barcode
            for field, label in (("cost_price", "preço de custo"), ("sale_price", "preço de venda")):
                try:
                    value = parse_number(get(r, field))
                except ValueError:
                    raise ValueError(f"{label} inválido") from None
                if value is not None:
                    if value < 0:
                        raise ValueError(f"{label} negativo")
                    data[field] = round(value, 2)
            for field, label in (("current_stock", "estoque"), ("min_stock", "estoque mínimo")):
                try:
                    value = parse_number(get(r, field))
                except ValueError:
                    raise ValueError(f"{label} inválido") from None
                if value is not None:
                    if value < 0:
                        raise ValueError(f"{label} negativo")
                    data[field] = int(value)

            current = existing.get(sku)
            if current is None:
                if not name:
                    raise ValueError("Nome vazio")
                if "sale_price" not in data:
                    raise ValueError("produto novo precisa do preço de venda")
                action = "CRIAR"
            else:
                changed = {k: v for k, v in data.items() if current.get(k) != v}
                data = changed
                action = "ATUALIZAR" if changed else "SEM_MUDANCA"
            if barcode:
                barcode_owner[barcode] = sku
            plans.append((action, {**data, "sku": sku}, current))
            rows.append(ImportRow(row=line, sku=sku, name=name or (current or {}).get("name", ""), action=action))
        except ValueError as exc:
            rows.append(ImportRow(row=line, sku=sku, name=name, action="ERRO", message=str(exc)))

    created = sum(1 for a, _, _ in plans if a == "CRIAR")
    updated = sum(1 for a, _, _ in plans if a == "ATUALIZAR")
    unchanged = sum(1 for a, _, _ in plans if a == "SEM_MUDANCA")
    errors = sum(1 for r in rows if r.action == "ERRO")

    blocked = ""
    org = await db.organizations.find_one({"id": principal.org_id}) or {}
    limit = (await limits_for_org(org)).get("products")
    if limit is not None and created and len(existing) + created > limit:
        plan = PLAN_LABELS.get(org.get("plan", ""), org.get("plan", ""))
        blocked = (
            f"Seu plano {plan} permite até {limit} produtos e você já tem {len(existing)}. "
            f"A planilha cria {created} novos. Mude de plano em Minha assinatura ou importe menos produtos."
        )

    def result(applied: bool) -> ImportResult:
        verb = "Importação concluída" if applied else "Prévia"
        return ImportResult(
            dry_run=not applied,
            total_rows=len(body),
            created=created,
            updated=updated,
            unchanged=unchanged,
            errors=errors,
            rows=rows[:PREVIEW_ROWS] if not applied else [r for r in rows if r.action == "ERRO"][:PREVIEW_ROWS],
            blocked=blocked,
            message=f"{verb}: {created} novos, {updated} atualizados, {unchanged} sem mudança, {errors} com erro.",
        )

    if dry_run or blocked:
        if blocked and not dry_run:
            raise HTTPException(status_code=403, detail=blocked, headers={"X-Plan-Limit": "products"})
        return result(applied=False)

    now = datetime.now(timezone.utc)
    reason = "Importação de planilha"
    new_products: list[dict] = []
    product_updates: list[UpdateOne] = []
    movements: list[dict] = []
    for action, data, current in plans:
        if action == "SEM_MUDANCA":
            continue
        if action == "CRIAR":
            product = Product(
                org_id=principal.org_id,
                name=data["name"],
                sku=data["sku"],
                category=data.get("category", DEFAULT_CATEGORY),
                cost_price=data.get("cost_price", 0.0),
                sale_price=data["sale_price"],
                current_stock=data.get("current_stock", 0),
                min_stock=data.get("min_stock", 0),
                barcode=data.get("barcode", ""),
            )
            new_products.append(product.model_dump())
            if product.current_stock > 0:
                movements.append(
                    StockMovement(
                        org_id=principal.org_id, product_id=product.id, product_name=product.name,
                        sku=product.sku, type="ENTRADA", quantity=product.current_stock,
                        delta=product.current_stock, stock_after=product.current_stock,
                        reason=reason, operator_name=principal.name,
                    ).model_dump()
                )
            continue
        changes = {k: v for k, v in data.items() if k != "sku"}
        changes["updated_at"] = now
        product_updates.append(UpdateOne(scoped(principal, {"id": current["id"]}), {"$set": changes}))
        if "current_stock" in changes:
            delta = changes["current_stock"] - int(current.get("current_stock", 0))
            movements.append(
                StockMovement(
                    org_id=principal.org_id, product_id=current["id"],
                    product_name=changes.get("name", current["name"]), sku=current["sku"],
                    type="AJUSTE_POSITIVO" if delta > 0 else "AJUSTE_NEGATIVO",
                    quantity=abs(delta), delta=delta, stock_after=changes["current_stock"],
                    reason=reason, operator_name=principal.name,
                ).model_dump()
            )
    # grava em lotes: 50 mil linhas uma a uma estouraria o tempo da requisição
    for i in range(0, len(new_products), BATCH):
        await db.products.insert_many(new_products[i : i + BATCH], ordered=False)
    for i in range(0, len(product_updates), BATCH):
        await db.products.bulk_write(product_updates[i : i + BATCH], ordered=False)
    for i in range(0, len(movements), BATCH):
        await db.stock_movements.insert_many(movements[i : i + BATCH], ordered=False)
    return result(applied=True)
