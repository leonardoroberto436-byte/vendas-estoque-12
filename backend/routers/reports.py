"""Monthly closing report — mounted under /api/reports."""

from calendar import monthrange
from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse

from lib.auth import Principal, require_org, scoped

from lib.db import db
from models.inventory import (
    MonthlyReport,
    PaymentBreakdown,
    ReportDayPoint,
    ReportMonth,
    TopProduct,
)

router = APIRouter(prefix="/reports", tags=["reports"])

BR_TZ = ZoneInfo("America/Sao_Paulo")
MONTH_NAMES = [
    "janeiro",
    "fevereiro",
    "março",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


def _local(value) -> datetime:
    if isinstance(value, datetime) and value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(BR_TZ)


@router.get("/months", response_model=list[ReportMonth])
async def available_months(principal: Principal = Depends(require_org)):
    """Months that already have at least one sale, newest first (plus the current month)."""
    docs = await db.sales.find(scoped(principal), {"created_at": 1}).to_list(20000)
    keys: set[tuple[int, int]] = set()
    for doc in docs:
        local = _local(doc["created_at"])
        keys.add((local.year, local.month))
    now = datetime.now(timezone.utc).astimezone(BR_TZ)
    keys.add((now.year, now.month))
    return [
        ReportMonth(year=year, month=month, label=f"{MONTH_NAMES[month - 1].capitalize()} {year}")
        for year, month in sorted(keys, reverse=True)
    ]


@router.get("/monthly", response_model=MonthlyReport)
async def monthly_report(
    year: int | None = Query(default=None, ge=2000, le=2100),
    month: int | None = Query(default=None, ge=1, le=12),
    principal: Principal = Depends(require_org),
):
    now = datetime.now(timezone.utc).astimezone(BR_TZ)
    year = year or now.year
    month = month or now.month

    days_in_month = monthrange(year, month)[1]
    start = datetime(year, month, 1, tzinfo=BR_TZ).astimezone(timezone.utc)
    end_year, end_month = (year + 1, 1) if month == 12 else (year, month + 1)
    end = datetime(end_year, end_month, 1, tzinfo=BR_TZ).astimezone(timezone.utc)

    sales = await db.sales.find(
        scoped(principal, {"created_at": {"$gte": start, "$lt": end}})
    ).to_list(20000)

    revenue = round(sum(float(s["total"]) for s in sales), 2)
    cost = round(sum(float(s.get("cost_total", 0)) for s in sales), 2)
    profit = round(sum(float(s.get("profit", 0)) for s in sales), 2)
    discount = round(sum(float(s.get("discount", 0)) for s in sales), 2)
    items_sold = sum(int(i["quantity"]) for s in sales for i in s.get("items", []))

    daily = {day: {"revenue": 0.0, "profit": 0.0, "sales_count": 0} for day in range(1, days_in_month + 1)}
    payments: dict[str, dict] = {}
    tally: dict[str, dict] = {}

    for sale in sales:
        day = _local(sale["created_at"]).day
        bucket = daily[day]
        bucket["revenue"] += float(sale["total"])
        bucket["profit"] += float(sale.get("profit", 0))
        bucket["sales_count"] += 1

        pay = payments.setdefault(sale["payment_method"], {"revenue": 0.0, "sales_count": 0})
        pay["revenue"] += float(sale["total"])
        pay["sales_count"] += 1

        for item in sale.get("items", []):
            entry = tally.setdefault(
                item["product_id"],
                {"name": item["product_name"], "quantity": 0, "revenue": 0.0},
            )
            entry["quantity"] += int(item["quantity"])
            entry["revenue"] += float(item["line_total"])

    best_day = max(daily.items(), key=lambda kv: kv[1]["revenue"]) if sales else None
    top = sorted(tally.items(), key=lambda kv: kv[1]["revenue"], reverse=True)[:10]

    return MonthlyReport(
        year=year,
        month=month,
        label=f"{MONTH_NAMES[month - 1].capitalize()} {year}",
        generated_at=datetime.now(timezone.utc),
        sales_count=len(sales),
        items_sold=items_sold,
        revenue=revenue,
        cost=cost,
        profit=profit,
        discount_total=discount,
        margin_percent=round((profit / revenue * 100) if revenue else 0, 1),
        average_ticket=round(revenue / len(sales), 2) if sales else 0.0,
        best_day=best_day[0] if best_day and best_day[1]["revenue"] > 0 else None,
        best_day_revenue=round(best_day[1]["revenue"], 2) if best_day else 0.0,
        daily=[
            ReportDayPoint(
                day=day,
                revenue=round(value["revenue"], 2),
                profit=round(value["profit"], 2),
                sales_count=value["sales_count"],
            )
            for day, value in daily.items()
        ],
        payments=[
            PaymentBreakdown(
                payment_method=method,
                revenue=round(value["revenue"], 2),
                sales_count=value["sales_count"],
                share_percent=round((value["revenue"] / revenue * 100) if revenue else 0, 1),
            )
            for method, value in sorted(payments.items(), key=lambda kv: -kv[1]["revenue"])
        ],
        top_products=[
            TopProduct(
                product_id=pid,
                name=data["name"],
                quantity=data["quantity"],
                revenue=round(data["revenue"], 2),
            )
            for pid, data in top
        ],
    )


@router.get("/monthly.xlsx")
async def monthly_report_xlsx(
    year: int | None = Query(default=None, ge=2000, le=2100),
    month: int | None = Query(default=None, ge=1, le=12),
    principal: Principal = Depends(require_org),
):
    """Same closing figures as /monthly, as a four-sheet .xlsx for the accountant."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    report = await monthly_report(year=year, month=month, principal=principal)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D9488")
    title_font = Font(bold=True, size=14)
    money = 'R$ #,##0.00'

    wb = Workbook()

    def style_header(ws, row: int, columns: int) -> None:
        for col in range(1, columns + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

    def autosize(ws, widths: list[int]) -> None:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width

    # --- Resumo
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Estoque & Vendas Pro — Fechamento {report.label}"
    ws["A1"].font = title_font
    ws["A2"] = f"Emitido em {report.generated_at.astimezone(BR_TZ).strftime('%d/%m/%Y %H:%M')}"
    rows = [
        ("Faturamento", report.revenue, money),
        ("Custo das vendas", report.cost, money),
        ("Lucro", report.profit, money),
        ("Margem (%)", report.margin_percent, "0.0"),
        ("Descontos concedidos", report.discount_total, money),
        ("Vendas realizadas", report.sales_count, "0"),
        ("Itens vendidos", report.items_sold, "0"),
        ("Ticket médio", report.average_ticket, money),
        ("Melhor dia", report.best_day or "—", "General"),
        ("Faturamento do melhor dia", report.best_day_revenue, money),
    ]
    ws["A4"] = "Indicador"
    ws["B4"] = "Valor"
    style_header(ws, 4, 2)
    for offset, (label, value, fmt) in enumerate(rows, start=5):
        ws.cell(row=offset, column=1, value=label)
        cell = ws.cell(row=offset, column=2, value=value)
        cell.number_format = fmt
    autosize(ws, [30, 22])

    # --- Produtos mais vendidos
    ws = wb.create_sheet("Produtos")
    ws.append(["#", "Produto", "Quantidade", "Faturamento"])
    style_header(ws, 1, 4)
    for index, item in enumerate(report.top_products, start=1):
        ws.append([index, item.name, item.quantity, item.revenue])
        ws.cell(row=index + 1, column=4).number_format = money
    autosize(ws, [6, 40, 14, 18])

    # --- Formas de pagamento
    ws = wb.create_sheet("Pagamentos")
    labels = {
        "DINHEIRO": "Dinheiro",
        "PIX": "PIX",
        "CARTAO_DEBITO": "Cartão de Débito",
        "CARTAO_CREDITO": "Cartão de Crédito",
    }
    ws.append(["Forma de pagamento", "Vendas", "Faturamento", "Participação (%)"])
    style_header(ws, 1, 4)
    for index, row in enumerate(report.payments, start=2):
        ws.append(
            [
                labels.get(row.payment_method, row.payment_method),
                row.sales_count,
                row.revenue,
                row.share_percent,
            ]
        )
        ws.cell(row=index, column=3).number_format = money
        ws.cell(row=index, column=4).number_format = "0.0"
    autosize(ws, [24, 12, 18, 18])

    # --- Faturamento por dia
    ws = wb.create_sheet("Por dia")
    ws.append(["Dia", "Vendas", "Faturamento", "Lucro"])
    style_header(ws, 1, 4)
    for index, point in enumerate(report.daily, start=2):
        ws.append([point.day, point.sales_count, point.revenue, point.profit])
        ws.cell(row=index, column=3).number_format = money
        ws.cell(row=index, column=4).number_format = money
    autosize(ws, [8, 12, 18, 18])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"fechamento-{report.year}-{report.month:02d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
