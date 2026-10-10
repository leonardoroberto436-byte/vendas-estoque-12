"""Company branding — mounted under /api/settings.

Tenant-scoped: the branding always belongs to the caller's own organization
(`principal.org_id`), so one company can never read or overwrite another's logo.
Reading is open to any member (the POS needs it); writing is ADMIN only.
"""

from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from lib.auth import Principal, require, require_org, scoped
from lib.db import db
from models.branding import LOGO_MAX_CHARS, AlertSettings, OrgBranding, OrgBrandingUpdate

BR_TZ = ZoneInfo("America/Sao_Paulo")

router = APIRouter(prefix="/settings", tags=["settings"])

_ALLOWED_LOGO_PREFIXES = ("data:image/png;base64,", "data:image/jpeg;base64,")


def _branding(org: dict | None) -> OrgBranding:
    data = (org or {}).get("branding") or {}
    return OrgBranding(
        trade_name=data.get("trade_name") or (org or {}).get("name", ""),
        cnpj=data.get("cnpj", ""),
        phone=data.get("phone", ""),
        address=data.get("address", ""),
        logo=data.get("logo"),
    )


@router.get("/branding", response_model=OrgBranding)
async def get_branding(principal: Principal = Depends(require_org)):
    org = await db.organizations.find_one({"id": principal.org_id})
    return _branding(org)


@router.put("/branding", response_model=OrgBranding)
async def update_branding(
    payload: OrgBrandingUpdate, principal: Principal = Depends(require("users:manage"))
):
    if not principal.org_id:
        raise HTTPException(status_code=403, detail="Usuário sem empresa vinculada.")

    logo = payload.logo
    if logo:
        if not logo.startswith(_ALLOWED_LOGO_PREFIXES):
            raise HTTPException(status_code=400, detail="O logo deve ser uma imagem PNG ou JPG.")
        if len(logo) > LOGO_MAX_CHARS:
            raise HTTPException(status_code=413, detail="O logo deve ter no máximo 300 KB.")

    await db.organizations.update_one(
        {"id": principal.org_id},
        {
            "$set": {
                "branding": {
                    "trade_name": payload.trade_name.strip(),
                    "cnpj": payload.cnpj.strip(),
                    "phone": payload.phone.strip(),
                    "address": payload.address.strip(),
                    "logo": logo,
                }
            }
        },
    )
    org = await db.organizations.find_one({"id": principal.org_id})
    return _branding(org)


# ------------------------------------------------------------ alertas de estoque
async def low_stock_alerts_enabled(org_id: str | None) -> bool:
    """Alertas ligados por padrão; a empresa pode desligar em Alertas."""
    if not org_id:
        return True
    org = await db.organizations.find_one({"id": org_id}, {"alerts": 1})
    return bool(((org or {}).get("alerts") or {}).get("low_stock_enabled", True))


@router.get("/alerts", response_model=AlertSettings)
async def get_alert_settings(principal: Principal = Depends(require_org)):
    return AlertSettings(low_stock_enabled=await low_stock_alerts_enabled(principal.org_id))


@router.put("/alerts", response_model=AlertSettings)
async def update_alert_settings(
    payload: AlertSettings, principal: Principal = Depends(require("users:manage"))
):
    """Só o administrador da loja liga ou desliga os alertas."""
    if not principal.org_id:
        raise HTTPException(status_code=403, detail="Usuário sem empresa vinculada.")
    await db.organizations.update_one(
        {"id": principal.org_id},
        {"$set": {"alerts.low_stock_enabled": payload.low_stock_enabled}},
    )
    return AlertSettings(low_stock_enabled=payload.low_stock_enabled)


# ------------------------------------------------------ exportação dos dados
def _local(value) -> str:
    if not isinstance(value, datetime):
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(BR_TZ).strftime("%d/%m/%Y %H:%M")


@router.get("/export.xlsx")
async def export_company_data(principal: Principal = Depends(require("users:manage"))):
    """Todos os dados da empresa numa planilha — cópia que a própria loja guarda.

    Produtos, vendas (com itens), devoluções e movimentações de estoque. Só o
    administrador baixa, e só os dados da própria empresa.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    products = await db.products.find(scoped(principal)).sort("name", 1).to_list(50000)
    sales = await db.sales.find(scoped(principal)).sort("created_at", 1).to_list(200000)
    returns = await db.returns.find(scoped(principal)).sort("created_at", 1).to_list(200000)
    movements = await db.stock_movements.find(scoped(principal)).sort("created_at", 1).to_list(500000)

    wb = Workbook()
    bold = Font(bold=True)

    def sheet(title: str, header: list[str], rows: list[list], first: bool = False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"

    sheet(
        "Produtos",
        ["SKU", "Produto", "Categoria", "Custo", "Preço de venda", "Estoque", "Estoque mínimo"],
        [
            [p.get("sku"), p.get("name"), p.get("category"), p.get("cost_price"),
             p.get("sale_price"), p.get("current_stock"), p.get("min_stock")]
            for p in products
        ],
        first=True,
    )
    sheet(
        "Vendas",
        ["Código", "Data", "Vendedor", "Pagamento", "Subtotal", "Desconto", "Total",
         "Custo", "Lucro", "Devolvido", "Observação"],
        [
            [s.get("code"), _local(s.get("created_at")), s.get("seller_name"), s.get("payment_method"),
             s.get("subtotal"), s.get("discount"), s.get("total"), s.get("cost_total"),
             s.get("profit"), s.get("returned_total", 0), s.get("notes", "")]
            for s in sales
        ],
    )
    sheet(
        "Itens vendidos",
        ["Venda", "Data", "SKU", "Produto", "Quantidade", "Preço unitário", "Total da linha"],
        [
            [s.get("code"), _local(s.get("created_at")), i.get("sku"), i.get("product_name"),
             i.get("quantity"), i.get("unit_price"), i.get("line_total")]
            for s in sales
            for i in s.get("items", [])
        ],
    )
    sheet(
        "Devoluções",
        ["Venda", "Data", "Operador", "Itens", "Valor devolvido", "Voltou ao estoque", "Motivo"],
        [
            [r.get("sale_code"), _local(r.get("created_at")), r.get("operator_name"),
             ", ".join(f"{i.get('quantity')}x {i.get('product_name')}" for i in r.get("items", [])),
             r.get("refund_total"), "sim" if r.get("restocked") else "não", r.get("reason", "")]
            for r in returns
        ],
    )
    sheet(
        "Movimentações",
        ["Data", "SKU", "Produto", "Tipo", "Quantidade", "Variação", "Estoque depois",
         "Responsável", "Motivo"],
        [
            [_local(m.get("created_at")), m.get("sku"), m.get("product_name"), m.get("type"),
             m.get("quantity"), m.get("delta"), m.get("stock_after"), m.get("operator_name"),
             m.get("reason", "")]
            for m in movements
        ],
    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    stamp = datetime.now(timezone.utc).astimezone(BR_TZ).strftime("%Y-%m-%d")
    filename = f"dados-da-loja-{stamp}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
