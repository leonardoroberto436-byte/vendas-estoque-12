"""Código de barras nos produtos + importação de produtos por planilha.

Run: cd /app/backend && python -m pytest tests/test_import_barcode.py -q
"""

import io
import os
import uuid
from contextlib import contextmanager

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from tests.credentials import env

API_URL = os.environ.get(
    "BACKEND_URL", "https://vendas-estoque-12.preview.emergentagent.com"
) + "/api"

PASSWORD = "Empresa@2026"
PLATFORM_EMAIL = env("PLATFORM_OWNER_EMAIL")
PLATFORM_PASSWORD = env("PLATFORM_OWNER_PASSWORD")


@contextmanager
def _session(email: str, password: str):
    c = httpx.Client(base_url=API_URL, timeout=60.0)
    resp = c.post("/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, resp.text
    try:
        yield c
    finally:
        c.close()


@pytest.fixture
def company():
    email = f"tscheck-imp.{uuid.uuid4().hex[:8]}@teste-saas.com"
    name = f"TscheckImp {uuid.uuid4().hex[:6]}"
    with httpx.Client(base_url=API_URL, timeout=60.0) as c:
        resp = c.post(
            "/auth/signup",
            json={"name": "Dono", "company_name": name, "email": email, "password": PASSWORD, "accept_terms": True},
        )
    assert resp.status_code == 201, resp.text
    info = {"email": email, "org_id": resp.json()["organization_id"], "name": name}
    yield info
    with _session(PLATFORM_EMAIL, PLATFORM_PASSWORD) as owner:
        owner.request("DELETE", f"/platform/organizations/{info['org_id']}", json={"confirm_name": name})


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(admin: httpx.Client, content: bytes, name: str = "produtos.xlsx", dry_run: bool = True):
    return admin.post(
        "/products/import",
        params={"dry_run": str(dry_run).lower()},
        files={"file": (name, content)},
    )


# ------------------------------------------------------------ código de barras
def test_barcode_is_saved_searchable_and_unique(company):
    with _session(company["email"], PASSWORD) as admin:
        base = {"category": "Freios", "cost_price": 10, "sale_price": 20, "current_stock": 3, "min_stock": 0}
        a = admin.post("/products", json={**base, "name": "Pastilha", "sku": "PA-1", "barcode": " 7891234567895 "})
        assert a.status_code == 201, a.text
        assert a.json()["barcode"] == "7891234567895"

        dup = admin.post("/products", json={**base, "name": "Outra", "sku": "PA-2", "barcode": "7891234567895"})
        assert dup.status_code == 409, dup.text

        # sem código de barras é permitido em vários produtos
        for sku in ("PA-3", "PA-4"):
            assert admin.post("/products", json={**base, "name": sku, "sku": sku}).status_code == 201

        found = admin.get("/products", params={"search": "7891234567895"}).json()
        assert [p["sku"] for p in found] == ["PA-1"]

        b = admin.get("/products", params={"search": "PA-3"}).json()[0]
        clash = admin.put(f"/products/{b['id']}", json={"barcode": "7891234567895"})
        assert clash.status_code == 409, clash.text
        ok = admin.put(f"/products/{a.json()['id']}", json={"barcode": "7891234567895", "name": "Pastilha X"})
        assert ok.status_code == 200, ok.text


# ------------------------------------------------------------ importação
def test_template_downloads(company):
    with _session(company["email"], PASSWORD) as admin:
        resp = admin.get("/products/import/template.xlsx")
        assert resp.status_code == 200, resp.text
        ws = load_workbook(io.BytesIO(resp.content)).active
        assert [c.value for c in ws[1]][:2] == ["Nome", "SKU"]


def test_import_preview_then_apply_creates_and_updates(company):
    with _session(company["email"], PASSWORD) as admin:
        admin.post(
            "/products",
            json={"name": "Filtro", "sku": "FO-1", "category": "Filtros", "cost_price": 5,
                  "sale_price": 10, "current_stock": 4, "min_stock": 1},
        )
        content = _xlsx([
            ["Lista de preços"],  # título acima da tabela
            ["Descrição", "Código", "EAN", "Grupo", "Custo", "Preço de venda", "Qtd", "Estoque mínimo"],
            ["Pastilha", "pf-10", "7890000000017", "Freios", "45,90", "R$ 1.089,90", 12, 3],
            ["Filtro", "FO-1", None, None, None, "12,50", 9, None],  # atualiza preço e estoque
            ["Sem preço", "NP-1", None, None, None, None, 1, None],  # erro: novo sem preço
            ["Repetido", "PF-10", None, None, None, "5", None, None],  # erro: SKU repetido
            [None, None, None, None, None, None, None, None],  # linha vazia ignorada
        ])
        preview = _upload(admin, content)
        assert preview.status_code == 200, preview.text
        body = preview.json()
        assert body["dry_run"] is True
        assert (body["created"], body["updated"], body["errors"]) == (1, 1, 2)
        assert {r["row"]: r["action"] for r in body["rows"]} == {3: "CRIAR", 4: "ATUALIZAR", 5: "ERRO", 6: "ERRO"}
        assert len(admin.get("/products").json()) == 1  # nada gravado na prévia

        applied = _upload(admin, content, dry_run=False)
        assert applied.status_code == 200, applied.text
        assert applied.json()["created"] == 1

        products = {p["sku"]: p for p in admin.get("/products").json()}
        assert products["PF-10"]["sale_price"] == 1089.90
        assert products["PF-10"]["cost_price"] == 45.90
        assert products["PF-10"]["barcode"] == "7890000000017"
        assert products["PF-10"]["current_stock"] == 12
        assert products["FO-1"]["sale_price"] == 12.50
        assert products["FO-1"]["current_stock"] == 9
        assert products["FO-1"]["category"] == "Filtros"  # coluna vazia não apaga

        movements = admin.get("/stock/movements").json()
        reasons = [m for m in movements if m.get("reason") == "Importação de planilha"]
        assert len(reasons) == 2

        again = _upload(admin, content).json()
        assert again["created"] == 0 and again["unchanged"] == 2


def test_import_csv_semicolon(company):
    csv_text = "Nome;SKU;Categoria;Preço;Estoque\nÓleo 5W30;OL-1;Óleos;39,90;7\n"
    with _session(company["email"], PASSWORD) as admin:
        resp = _upload(admin, csv_text.encode("cp1252"), name="lista.csv", dry_run=False)
        assert resp.status_code == 200, resp.text
        product = admin.get("/products").json()[0]
        assert (product["name"], product["sale_price"], product["current_stock"]) == ("Óleo 5W30", 39.9, 7)


def test_import_respects_plan_limit(company):
    with _session(PLATFORM_EMAIL, PLATFORM_PASSWORD) as owner:
        resp = owner.put(
            f"/platform/organizations/{company['org_id']}/plan",
            json={"plan": "INICIAL", "subscription_status": "ATIVO", "due_date": None},
        )
        assert resp.status_code == 200, resp.text
        limit = owner.get("/platform/plan-limits").json()["limits"]["INICIAL"]["products"]
    rows = [["Nome", "SKU", "Preço"]] + [[f"P{i}", f"S{i}", 1] for i in range(limit + 1)]
    with _session(company["email"], PASSWORD) as admin:
        preview = _upload(admin, _xlsx(rows)).json()
        assert preview["blocked"]
        refused = _upload(admin, _xlsx(rows), dry_run=False)
        assert refused.status_code == 403
        assert admin.get("/products").json() == []


def test_import_without_name_and_sku_columns_is_refused(company):
    with _session(company["email"], PASSWORD) as admin:
        resp = _upload(admin, _xlsx([["Coisa", "Outra"], [1, 2]]))
        assert resp.status_code == 400
        assert "Nome e SKU" in resp.json()["detail"]


def test_operator_cannot_import(company):
    with _session(company["email"], PASSWORD) as admin:
        op_email = f"tscheck-imp-op.{uuid.uuid4().hex[:6]}@teste-saas.com"
        created = admin.post(
            "/users", json={"name": "Op", "email": op_email, "password": "Operador@2026", "role": "OPERADOR"}
        )
        assert created.status_code == 201, created.text
    with _session(op_email, "Operador@2026") as op:
        assert _upload(op, _xlsx([["Nome", "SKU", "Preço"], ["A", "B", 1]])).status_code == 403


# ------------------------------------------------------------ página de vendas
def test_public_plans_without_login():
    with httpx.Client(base_url=API_URL, timeout=60.0) as c:
        resp = c.get("/public/plans")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["trial_days"] == 7
    assert [p["id"] for p in body["plans"]] == ["INICIAL", "PROFISSIONAL", "AVANCADO"]
    assert all(p["price"] > 0 for p in body["plans"])
